"""Generate the student-friendly, five-slide CrimeGraph AI project review."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent.parent
OUTPUT = ROOT / "CrimeGraph_AI_Presentation_Streamlit.pptx"

NAVY = "102033"
INK = "14213A"
CARD = "192B43"
WHITE = "FFFFFF"
MUTED = "536276"
PALE = "F3F6F9"
CYAN = "168A8A"
GOLD = "DA9A3C"
LINE = "D7DFE8"
GREEN = "477F62"


def rgb(value):
    return RGBColor.from_string(value)


def add_text(slide, value, x, y, w, h, size=14, color=INK, bold=False, align=PP_ALIGN.LEFT):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.05)
    frame.margin_right = Inches(0.05)
    frame.margin_top = Inches(0.035)
    frame.margin_bottom = Inches(0.025)
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    for index, line in enumerate(value.split("\n")):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = line
        paragraph.alignment = align
        paragraph.font.name = "Aptos"
        paragraph.font.size = Pt(size)
        paragraph.font.bold = bold
        paragraph.font.color.rgb = rgb(color)
    return shape


def add_box(slide, x, y, w, h, fill=PALE, outline=LINE, radius=True):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        Inches(x),
        Inches(y),
        Inches(w),
        Inches(h),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.line.color.rgb = rgb(outline)
    shape.line.width = Pt(1)
    return shape


def add_header(slide, number, label, title, subtitle):
    add_text(slide, f"CRIMEGRAPH AI     /     PROJECT REVIEW                                      {number} / 5", 0.62, 0.25, 12.0, 0.2, 8, CYAN, True)
    add_text(slide, label.upper(), 0.68, 0.78, 11.9, 0.24, 10, GOLD, True)
    add_text(slide, title, 0.62, 1.1, 12.0, 0.58, 27, NAVY, True)
    add_text(slide, subtitle, 0.67, 1.76, 11.9, 0.36, 12.5, MUTED)
    add_box(slide, 0.62, 7.03, 12.05, 0.012, LINE, LINE, radius=False)
    add_text(slide, "CLASSROOM PROTOTYPE   •   FICTIONAL / GENERATED DATA   •   HUMAN REVIEW REQUIRED", 0.65, 7.1, 11.4, 0.18, 7.8, MUTED)


def new_slide(pres, number, label, title, subtitle):
    slide = pres.slides.add_slide(pres.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb("FFFFFF")
    add_header(slide, number, label, title, subtitle)
    return slide


def add_arrow(slide, x1, x2, y):
    line = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Inches(x1),
        Inches(y),
        Inches(x2),
        Inches(y),
    )
    line.line.color.rgb = rgb(CYAN)
    line.line.width = Pt(2.4)
    line.line.end_arrowhead = True


def add_case_graph(slide):
    nodes = {
        "john": (7.56, 3.34, 0.98, 0.4, "John Mercer", "B84F59"),
        "marcus": (7.56, 4.45, 1.0, 0.4, "Marcus Vance", "B84F59"),
        "elena": (7.56, 5.34, 1.0, 0.4, "Elena Rostova", "B84F59"),
        "phone1": (9.08, 2.96, 0.84, 0.4, "Phone 01", "467FAF"),
        "phone2": (9.04, 5.34, 0.88, 0.4, "Phone 02", "467FAF"),
        "location": (8.38, 3.84, 1.12, 0.4, "Central Station", "477F62"),
        "case": (9.72, 4.04, 1.28, 0.52, "Blackout Case", CYAN),
        "file": (10.92, 3.04, 0.9, 0.4, "Payload file", "9867A8"),
        "ip_internal": (10.8, 5.02, 1.12, 0.4, "192.168.1.105", "7967B6"),
        "ip_external": (11.2, 4.12, 1.14, 0.4, "198.51.100.22", "7967B6"),
    }
    centers = {
        key: (x + width / 2, y + height / 2)
        for key, (x, y, width, height, _, _) in nodes.items()
    }
    relationships = [
        ("john", "case"), ("marcus", "case"), ("elena", "case"),
        ("john", "phone1"), ("marcus", "phone2"), ("phone1", "phone2"),
        ("john", "location"), ("marcus", "location"), ("marcus", "ip_internal"),
        ("ip_internal", "ip_external"), ("marcus", "file"),
    ]
    add_box(slide, 7.05, 2.7, 5.35, 3.03, NAVY, NAVY)
    for source, target in relationships:
        x1, y1 = centers[source]
        x2, y2 = centers[target]
        connector = slide.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT,
            Inches(x1), Inches(y1), Inches(x2), Inches(y2),
        )
        connector.line.color.rgb = rgb("70829C")
        connector.line.width = Pt(1.15)
    for x, y, width, height, label, fill in nodes.values():
        add_box(slide, x, y, width, height, fill, fill)
        add_text(slide, label, x + 0.025, y + 0.035, width - 0.05, height - 0.07, 7.2, WHITE, True, PP_ALIGN.CENTER)


def build():
    pres = Presentation()
    pres.slide_width = Inches(13.333)
    pres.slide_height = Inches(7.5)

    # 1. Problem statement
    slide = new_slide(
        pres,
        "01",
        "Problem statement",
        "Digital case clues are hard to see together",
        "Notes, network records, and file samples often arrive in different formats.",
    )
    cards = [
        ("01", "Many input types", "Notes, network traffic, and file fragments are separate kinds of information.", CYAN),
        ("02", "Connections get missed", "People, devices, places, and evidence can be difficult to connect by eye.", GOLD),
        ("03", "Need one clear view", "A reviewer needs an organized case view to explore possible links.", GREEN),
    ]
    for index, (number, title, body, accent) in enumerate(cards):
        x = 0.75 + index * 4.18
        add_box(slide, x, 2.65, 3.75, 2.14, PALE, LINE)
        add_text(slide, number, x + 0.22, 2.87, 0.68, 0.38, 19, accent, True)
        add_text(slide, title, x + 0.22, 3.38, 3.23, 0.37, 17, NAVY, True)
        add_text(slide, body, x + 0.22, 3.9, 3.2, 0.66, 12.3, MUTED)
    add_box(slide, 0.78, 5.37, 11.78, 0.76, "FFF7E9", "F0D9B4")
    add_text(slide, "Goal: organize sample information and show possible links — not decide guilt.", 1.0, 5.56, 11.35, 0.34, 15, NAVY, True, PP_ALIGN.CENTER)
    slide.notes_slide.notes_text_frame.text = (
        "Problem: different types of sample information are hard to review together. "
        "The project aims to organize them and show possible links for a person to inspect. "
        "It does not make decisions about people."
    )

    # 2. Solution
    slide = new_slide(
        pres,
        "02",
        "Solution",
        "Three ways to add sample information",
        "Each workflow adds a result to one case workspace and its relationship graph.",
    )
    steps = [
        ("CASE NOTES", "A simple text-pattern finder looks for supported details.", CYAN),
        ("NETWORK FLOW", "A trained model predicts a traffic category.", GOLD),
        ("FILE SAMPLE", "A trained model predicts one of six file types.", GREEN),
    ]
    for index, (title, body, accent) in enumerate(steps):
        x = 0.84 + index * 4.12
        add_box(slide, x, 2.58, 3.55, 1.62, PALE, LINE)
        add_box(slide, x + 0.2, 2.8, 0.14, 1.17, accent, accent, radius=False)
        add_text(slide, title, x + 0.5, 2.82, 2.8, 0.31, 12, accent, True)
        add_text(slide, body, x + 0.5, 3.26, 2.8, 0.64, 13, NAVY)
        if index < 2:
            add_arrow(slide, x + 3.62, x + 4.02, 3.39)
    add_box(slide, 1.82, 4.82, 9.68, 1.0, INK, INK)
    add_text(slide, "ONE CASE GRAPH", 2.12, 4.98, 2.0, 0.24, 11, "91D7D3", True)
    add_text(slide, "Select a case  •  view graph  •  inspect entity and link tables", 4.05, 4.94, 7.02, 0.38, 15, WHITE, True, PP_ALIGN.CENTER)
    add_text(slide, "Two trained classifiers + one separate rule-based notes feature", 2.2, 6.04, 9.0, 0.3, 12, MUTED, align=PP_ALIGN.CENTER)
    slide.notes_slide.notes_text_frame.text = (
        "The app accepts notes, traffic values, and small file-byte samples. "
        "Notes use a rule-based pattern extractor. Only the traffic and file workflows use trained classifiers. "
        "The findings can be reviewed in the case graph and entity/relationship tables."
    )

    # 3. Architecture
    slide = new_slide(
        pres,
        "03",
        "Architecture",
        "Streamlit runs the analysis locally",
        "The interface calls Python services directly; this demo does not send HTTP API requests.",
    )
    columns = [
        (0.66, 2.05, "1  ENTER", "Streamlit UI", "Case notes\nOne traffic flow\nHex byte sample", "E9F3F7", CYAN),
        (3.02, 2.0, "2  VALIDATE", "Pydantic", "Checks the input\nand validates\nthe values", "F1F4F8", GOLD),
        (5.34, 2.68, "3  ANALYZE", "Python services", "Traffic: CNN-LSTM\nFiles: TF-IDF + MLP\nNotes: local rules", "EAF4EF", GREEN),
        (8.34, 1.86, "4  SAVE", "SQLite", "Case records,\nentities, and\nrelationships", "F7F0E3", GOLD),
        (10.48, 2.13, "5  DISPLAY", "Streamlit graph", "Graphviz view\nEntity and link\ntables", "E9F3F7", CYAN),
    ]
    for x, w, step, heading, body, fill, accent in columns:
        add_box(slide, x, 2.78, w, 2.15, fill, LINE)
        add_text(slide, step, x + 0.12, 3.01, w - 0.24, 0.23, 9, accent, True, PP_ALIGN.CENTER)
        add_text(slide, heading, x + 0.12, 3.4, w - 0.24, 0.36, 16, NAVY, True, PP_ALIGN.CENTER)
        add_text(slide, body, x + 0.12, 3.87, w - 0.24, 0.78, 11, MUTED, align=PP_ALIGN.CENTER)
    for left, right in [(2.73, 2.98), (5.05, 5.30), (8.07, 8.30), (10.22, 10.45)]:
        add_arrow(slide, left, right, 3.86)
    add_box(slide, 0.95, 5.6, 11.45, 0.67, "FFF7E9", "F0D9B4")
    add_text(slide, "Local Python workflow • no HTTP API calls • results need human review", 1.18, 5.78, 11.0, 0.3, 13.5, NAVY, True, PP_ALIGN.CENTER)
    slide.notes_slide.notes_text_frame.text = (
        "For the Streamlit demo, the interface calls the project's Python services directly in the same local app. "
        "Pydantic validates submitted values. Traffic and file samples use the saved classifiers; notes use local text patterns. "
        "Case data and links are saved in SQLite and displayed with Streamlit's Graphviz visualization and data tables. "
        "This demo does not send frontend-to-backend HTTP API requests."
    )

    # 4. Model comparison table
    slide = new_slide(
        pres,
        "04",
        "Model comparison",
        "The two classifiers used in the app",
        "A quick comparison of what goes in, what each model does, and what comes out.",
    )
    x0 = 0.78
    widths = [2.13, 2.45, 4.1, 2.73]
    x_positions = [x0]
    for width in widths[:-1]:
        x_positions.append(x_positions[-1] + width)
    headers = ["MODEL", "SAMPLE INPUT", "WHAT IT DOES", "TRAINING SETUP"]
    rows = [
        ["Traffic", "One flow → 8 × 15\nmodel window", "CNN-LSTM predicts benign or malicious traffic", "10,000 generated flows\n50 epochs"],
        ["File type", "Hexadecimal byte sample", "TF-IDF + MLP predicts one of six file categories", "6,000 generated samples\n50 epochs"],
    ]
    for col, heading in enumerate(headers):
        add_box(slide, x_positions[col], 2.52, widths[col] - 0.04, 0.55, INK, INK, radius=False)
        add_text(slide, heading, x_positions[col] + 0.12, 2.66, widths[col] - 0.28, 0.22, 10, WHITE, True)
    for row_index, row in enumerate(rows):
        y = 3.15 + row_index * 1.13
        for col, value in enumerate(row):
            fill = PALE if row_index == 0 else "F9FBFC"
            add_box(slide, x_positions[col], y, widths[col] - 0.04, 0.97, fill, LINE, radius=False)
            add_text(slide, value, x_positions[col] + 0.12, y + 0.13, widths[col] - 0.28, 0.69, 11.1, NAVY if col == 0 else MUTED, col == 0)
    add_box(slide, 0.83, 5.78, 11.73, 0.69, "FFF7E9", "F0D9B4")
    add_text(slide, "Important: the examples are generated by the training code — the original named datasets were not used in the default runs.", 1.02, 5.95, 11.3, 0.34, 11.1, NAVY, True, PP_ALIGN.CENTER)
    slide.notes_slide.notes_text_frame.text = (
        "These are the two trained classifiers integrated into the application. "
        "The Streamlit form accepts one flow with 15 values; the inference service scales it and repeats it across the model's 8-step input window. "
        "The file model reads a hexadecimal sample. "
        "Both use 50 epochs in the current training scripts. The examples are generated; the original named datasets "
        "were not used in the standard training runs."
    )

    # 5. Results
    slide = new_slide(
        pres,
        "05",
        "Results",
        "What the prototype produced",
        "Held-out test scores on generated examples — not evidence of real-world accuracy.",
    )
    headers = ["MODEL", "ACCURACY", "F1 SCORE"]
    widths = [2.78, 1.55, 1.55]
    x_positions = [0.78, 3.6, 5.17]
    for col, heading in enumerate(headers):
        add_box(slide, x_positions[col], 2.42, widths[col] - 0.04, 0.48, INK, INK, radius=False)
        add_text(slide, heading, x_positions[col] + 0.08, 2.54, widths[col] - 0.20, 0.24, 9.5, WHITE, True, PP_ALIGN.CENTER if col else PP_ALIGN.LEFT)
    for row_index, row in enumerate([
        ["Traffic CNN-LSTM", "98.40%", "99.18%"],
        ["File TF-IDF + MLP", "99.92%", "99.92%"],
    ]):
        y = 2.98 + row_index * 0.66
        for col, value in enumerate(row):
            add_box(slide, x_positions[col], y, widths[col] - 0.04, 0.56, PALE, LINE, radius=False)
            add_text(slide, value, x_positions[col] + 0.08, y + 0.12, widths[col] - 0.2, 0.3, 10.5, NAVY, col == 0, PP_ALIGN.CENTER if col else PP_ALIGN.LEFT)
    add_box(slide, 0.78, 4.53, 5.95, 0.95, "FFF7E9", "F0D9B4")
    add_text(slide, "These scores describe generated test examples only. Real-world performance is unknown.", 1.03, 4.71, 5.45, 0.56, 12, NAVY, True, PP_ALIGN.CENTER)
    add_text(slide, "SAMPLE CASE GRAPH", 7.08, 2.32, 5.3, 0.27, 10, CYAN, True, PP_ALIGN.CENTER)
    add_case_graph(slide)
    add_text(slide, "Fictional seeded case • simplified Streamlit / Graphviz illustration", 7.05, 5.83, 5.36, 0.26, 8.8, MUTED, align=PP_ALIGN.CENTER)
    add_text(slide, "Next: test on suitable real datasets • independently validate • improve security and storage", 0.93, 6.38, 11.5, 0.32, 11.2, GREEN, True, PP_ALIGN.CENTER)
    slide.notes_slide.notes_text_frame.text = (
        "The metrics are test scores on generated benchmark examples. Traffic accuracy is 98.40 percent and F1 is 99.18 percent. "
        "File accuracy and F1 are 99.92 percent. Because the examples are generated, these numbers do not show how the models "
        "would perform on real forensic data. The graph is an illustration of the fictional seeded case: 10 entities and 11 links. "
        "Future work includes real-data evaluation, independent validation, and stronger security and storage."
    )

    pres.save(OUTPUT)
    reopened = Presentation(OUTPUT)
    if len(reopened.slides) != 5:
        raise RuntimeError(f"Expected 5 slides; wrote {len(reopened.slides)}")
    print(f"Created {OUTPUT} ({OUTPUT.stat().st_size} bytes; 5 slides)")


if __name__ == "__main__":
    build()
